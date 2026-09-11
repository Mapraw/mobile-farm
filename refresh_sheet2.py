"""Refresh Test > Sheet2 from the source workbook (yellow-highlight mapping).

Usage (double-click or command line):
    python refresh_sheet2.py

Reads:
    <this folder>/../บันทึกการเลี้ยง CA week21-30 Y69 Try_1.xlsx   (source, one sheet = one farm)
Writes:
    <this folder>/../test.xlsx  (Sheet2 rows 3+, one row per source sheet)

Mapping = Sheet2 row 2 cell addresses. Special case: 'B8-1' = B8 date minus 1 day
(Chick-in date = first record date B8 minus one day).
"""
import datetime
import glob
import os
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
PARENT = os.path.dirname(HERE)


def find_source():
    cands = glob.glob(os.path.join(PARENT, "*.xlsx"))
    cands = [c for c in cands if os.path.basename(c) != "test.xlsx" and not os.path.basename(c).startswith("~$")]
    if not cands:
        raise FileNotFoundError("Source workbook not found next to test.xlsx")
    return cands[0]


def main():
    src_path = find_source()
    dst_path = os.path.join(PARENT, "test.xlsx")
    print("Source:", src_path)
    src = load_workbook(src_path, data_only=True)
    dst = load_workbook(dst_path)
    ws2 = dst["Sheet2"]
    headers = [c.value for c in ws2[1]]
    refs = [c.value for c in ws2[2]]
    n = 0
    for ws in src.worksheets:
        row = []
        for ref in refs:
            if ref == "B8-1":
                v = ws["B8"].value
                if isinstance(v, datetime.datetime):
                    v = (v - datetime.timedelta(days=1)).date()
                row.append(v)
            else:
                row.append(ws[ref].value)
        for j, v in enumerate(row):
            ws2.cell(row=3 + n, column=1 + j, value=v)
        n += 1
    dst.save(dst_path)
    print(f"Done: wrote {n} farm rows to test.xlsx > Sheet2 (rows 3..{2 + n})")


if __name__ == "__main__":
    main()
